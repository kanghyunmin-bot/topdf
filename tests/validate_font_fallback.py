import hashlib,json,pathlib,subprocess,sys,os
from docx import Document
from docx.shared import Pt
from pypdf import PdfReader
root=pathlib.Path(__file__).resolve().parent.parent
app=pathlib.Path(sys.argv[1]).resolve()
exe=app/'TopDF.exe' if (app/'TopDF.exe').exists() else app/'Contents/MacOS/TopDF'
out=root/'build/font-fallback-tests'/app.parent.name
out.mkdir(parents=True,exist_ok=True)
results=[]
for family,label,expected in [('TopDF Missing Serif 123','missing','NanumGothic'),('Arial' if sys.platform=='win32' else 'Liberation Serif','available','Arial' if sys.platform=='win32' else 'LiberationSerif')]:
 doc=Document()
 p=doc.add_paragraph()
 run=p.add_run('FONT FAMILY CHECK 123 깔끔한 고딕체 확인')
 run.font.name=family;run.font.size=Pt(16)
 source=out/(label+'.docx');doc.save(source)
 target=out/(label+'.pdf')
 before=hashlib.sha256(source.read_bytes()).hexdigest()
 process=subprocess.run([str(exe),'--convert-test',str(source),str(target)],capture_output=True,text=True,timeout=60)
 assert hashlib.sha256(source.read_bytes()).hexdigest()==before
 assert process.returncode==0,(process.stdout,process.stderr)
 names=[]
 for page in PdfReader(target).pages:
  for font in page['/Resources']['/Font'].get_object().values():
   names.append(str(font.get_object()['/BaseFont']))
 assert any(expected in name for name in names),(label,names,process.stderr)
 results.append({'input_font':family,'pdf_fonts':names,'passed':True})
 print('PASS:',label,names,flush=True)
(out/'report.json').write_text(json.dumps(results,indent=2))
# Force the HWPX fixture to request an absent serif family.
import re,zipfile
source=out/'missing.hwpx'
with zipfile.ZipFile(root/'tests/sample.hwpx') as original,zipfile.ZipFile(source,'w') as rewritten:
 for entry in original.infolist():
  data=original.read(entry.filename)
  if entry.filename=='Contents/header.xml':
   data=re.sub(r'face="[^"]+"','face="TopDF Missing Serif 123"',data.decode()).encode()
  rewritten.writestr(entry,data)
process=subprocess.run([str(exe),'--convert-test',str(source),str(out/'missing-hwpx.pdf')],capture_output=True,text=True,timeout=60)
assert process.returncode==0,process.stderr
names=[str(font.get_object()['/BaseFont']) for page in PdfReader(out/'missing-hwpx.pdf').pages for font in page['/Resources']['/Font'].get_object().values()]
if sys.platform=='win32':
 # The pure Rust PDF backend intentionally names resources F0/F1, unlike CoreText.
 # Verify the resolver's exact input-font hashes and byte-identical app output.
 fonts=app/'Engines/LibreOffice/share/fonts/truetype'
 expected_hashes={hashlib.sha256((fonts/name).read_bytes()).hexdigest() for name in ['NanumGothic-Regular.ttf','NanumGothic-Bold.ttf']}
 direct=out/'missing-hwpx-direct.pdf';report=out/'hwp-resolver-report.json'
 environment=dict(os.environ,TOPDF_HWP_FALLBACK_FONT='NanumGothic')
 command=[str(app/'Engines/hwp.exe'),'render',str(source),'--output',str(direct),'--format','pdf','--report',str(report)]
 for directory in [fonts,pathlib.Path(os.environ['LOCALAPPDATA'])/'Microsoft/Windows/Fonts',pathlib.Path(os.environ['WINDIR'])/'Fonts']:
  command.extend(['--font-dir',str(directory)])
 process=subprocess.run(command,env=environment,capture_output=True,text=True,timeout=60)
 assert process.returncode==0,process.stderr
 resolved=json.loads(report.read_text())
 assert resolved['complete'] and resolved['font_resolution_complete'],resolved
 assert {font['resolved_sha256'] for font in resolved['fonts']}==expected_hashes,resolved['fonts']
 assert resolved['font_coverage']['missing']==0,resolved
 assert direct.read_bytes()==(out/'missing-hwpx.pdf').read_bytes(),'App PDF differs from the verified resolver output'
else:
 assert any('NanumGothic' in name for name in names),names
results.append({'format':'hwpx','input_font':'TopDF Missing Serif 123','pdf_fonts':names,'passed':True,'identity_check':'exact resolver SHA256 and byte-identical app PDF' if sys.platform=='win32' else 'CoreText embedded PDF font names'})
print('PASS: HWPX missing',names,flush=True)
(out/'report.json').write_text(json.dumps(results,indent=2))
