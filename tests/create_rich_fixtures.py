"""Create layout regression fixtures; these are test tools, not app dependencies."""
import pathlib
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, Alignment
from pptx import Presentation
from pptx.util import Inches as SlideInches, Pt as SlidePt
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import ChartData


def create(folder):
    folder = pathlib.Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    image = folder / 'graphic.png'
    canvas = Image.new('RGB', (600, 200), '#cde6ff')
    ImageDraw.Draw(canvas).rectangle((20, 20, 580, 180), outline='navy', width=5)
    canvas.save(image)
    doc = Document()
    doc.styles['Normal'].font.name = 'Liberation Serif'
    doc.styles['Normal'].font.size = Pt(12)
    doc.sections[0].header.paragraphs[0].text = 'LAYOUT TEST HEADER'
    doc.add_heading('Text, table and image', 0)
    doc.add_paragraph('Typography and wrapping. ' * 35)
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'
    table.cell(0, 0).merge(table.cell(0, 2)).text = 'MERGED CELLS'
    table.cell(1, 0).text = 'MULTILINE\nTABLE\nCELL'
    table.cell(1, 1).text = 'Long table text with automatic wrapping. ' * 6
    doc.add_picture(str(image), width=Inches(4))
    doc.add_page_break()
    doc.add_paragraph('EXPLICIT PAGE BREAK')
    doc.save(folder / 'rich.docx')
    book = Workbook()
    sheet = book.active
    sheet.append(['Category', 'Value', 'Formula'])
    for i in range(1, 8):
        sheet.append(['Long label ' + str(i), i * 12, '=B' + str(i + 1) + '*2'])
    sheet.merge_cells('A11:C11')
    sheet['A11'] = 'MERGED PRINT CELL'
    sheet.column_dimensions['A'].width = 24
    for row in sheet:
        for cell in row:
            cell.font = Font(name='Liberation Sans', size=11)
            cell.alignment = Alignment(wrap_text=True)
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=8), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=8))
    sheet.add_chart(chart, 'E2')
    sheet.print_area = 'A1:M18'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    book.save(folder / 'rich.xlsx')
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    box = slide.shapes.add_textbox(SlideInches(.5), SlideInches(.2), SlideInches(8), SlideInches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = 'Slide layout and chart'
    run.font.name = 'Liberation Sans'
    run.font.size = SlidePt(28)
    slide.shapes.add_picture(str(image), SlideInches(.5), SlideInches(1.3), width=SlideInches(4))
    data = ChartData()
    data.categories = ['A', 'B', 'C']
    data.add_series('Values', [10, 25, 15])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, SlideInches(.5), SlideInches(3), SlideInches(8), SlideInches(3), data)
    slides.save(folder / 'rich.pptx')
    return [folder / ('rich.' + ext) for ext in ('docx', 'xlsx', 'pptx')]
